"""Relatórios produzidos exclusivamente a partir do snapshot armazenado."""
import csv
from hashlib import sha256
from io import BytesIO, StringIO
import json
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import reportlab
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle

from .comparacao import METRICAS
from schemas.cenarios import Receita

FONT_DIR = Path(reportlab.__file__).parent / "fonts"
pdfmetrics.registerFont(TTFont("Fundeb", str(FONT_DIR / "Vera.ttf")))
pdfmetrics.registerFont(TTFont("FundebBold", str(FONT_DIR / "VeraBd.ttf")))


def seguro(valor):
    if isinstance(valor, str) and valor.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + valor
    return valor


def metadados(resultado):
    return {k: v for k, v in resultado.items() if k not in {"dados", "comparativos", "series"}}


def linhas_comparativo(resultado, comparacao):
    conjuntos = [(None, resultado["comparativos"])] if resultado["tipo"] == "cenario" else [
        (s["ano"], s["comparativos"]) for s in resultado["series"]]
    linhas = []
    entrada = resultado["entrada"]
    hipoteses = json.loads(json.dumps(entrada))
    objeto = hipoteses.get("cenario", hipoteses)
    ajustes = objeto.pop("ajustes", [])
    objeto["quantidade_ajustes"] = len(ajustes)
    objeto["ajustes_sha256"] = sha256(json.dumps(ajustes, sort_keys=True).encode()).hexdigest()
    # Os ajustes completos estão no XLSX/snapshot; não repetir milhares deles por linha do CSV.
    hipoteses_json = json.dumps(hipoteses, ensure_ascii=False, separators=(",", ":"))
    for ano, comparativos in conjuntos:
        for row in comparativos[comparacao]["linhas"]:
            saida = {"cenario_id": resultado["id"], "base_id": resultado["base"]["id"],
                     "base_sha256": resultado["base"]["fingerprint"], "exercicio_base": resultado["base"].get("ano_exercicio"),
                     "ano_trajetoria": ano, "comparacao": comparacao,
                     **{k: row[k] for k in ("ibge", "uf", "nome", "tipo")}}
            for indicador, valores in row["indicadores"].items():
                for medida, valor in valores.items():
                    saida[f"{indicador}_{medida}"] = valor
            saida.update({f"participacao_{k}": v for k, v in row["participacao"].items()})
            saida["denominador_participacao"] = resultado["recorte"]["denominador"]
            saida["status_base"] = resultado["base"]["status"]
            saida["aviso"] = resultado["aviso"]
            saida["hipoteses_json"] = hipoteses_json
            linhas.append(saida)
    return linhas


def csv_bytes(resultado, comparacao):
    rows = linhas_comparativo(resultado, comparacao)
    stream = StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=list(rows[0]), delimiter=";")
    writer.writeheader()
    for row in rows:
        writer.writerow({k: seguro(v) for k, v in row.items()})
    return stream.getvalue().encode("utf-8-sig")


def adicionar_aba(wb, nome, rows):
    sheet = wb.create_sheet(nome)
    if not rows:
        sheet.append(["Sem registros"])
        return
    headers = list(rows[0])
    sheet.append(headers)
    for row in rows:
        sheet.append([seguro(json.dumps(row.get(h), ensure_ascii=False) if isinstance(row.get(h), (dict, list)) else row.get(h)) for h in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="204D46")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 44
    for col, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(col)].width = min(46, max(16, len(header) * 0.8))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = '#,##0.00;[Red]-#,##0.00'
    return sheet


def xlsx_bytes(resultado, comparacao):
    wb = Workbook()
    wb.remove(wb.active)
    if resultado["tipo"] == "cenario":
        resumo = [{"comparacao": c, "indicador": METRICAS[m], **v}
                  for c, comp in resultado["comparativos"].items() for m, v in comp["totais"].items()]
    else:
        resumo = [{"ano": s["ano"], "comparacao": c, "indicador": METRICAS[m], **v}
                  for s in resultado["series"] for c, comp in s["comparativos"].items() for m, v in comp["totais"].items()]
    adicionar_aba(wb, "Resumo", resumo)
    adicionar_aba(wb, "Comparativo", linhas_comparativo(resultado, comparacao))
    for c in ("B-A", "C-A", "D-A", "D-C"):
        if c != comparacao:
            adicionar_aba(wb, c, linhas_comparativo(resultado, c))
    if resultado["tipo"] == "cenario":
        for c, dados in resultado["dados"].items():
            adicionar_aba(wb, "Dados " + c, dados)
        adicionar_aba(wb, "Ajustes", resultado["ajustes"])
    else:
        adicionar_aba(wb, "Trajetória", [{k: v for k, v in s.items() if k not in {"comparativos", "resumos_nacionais"}} for s in resultado["series"]])
    # Uma linha por campo evita o limite de 32.767 caracteres por célula do Excel.
    def achatar(obj, prefixo=""):
        if isinstance(obj, dict):
            for k, v in obj.items():
                yield from achatar(v, f"{prefixo}.{k}" if prefixo else k)
        elif isinstance(obj, list):
            for i, v in enumerate(obj):
                yield from achatar(v, f"{prefixo}[{i}]")
        else:
            yield {"campo": prefixo, "valor": obj}
    adicionar_aba(wb, "Metadados", list(achatar(metadados(resultado))))
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def numero(valor):
    if valor is None:
        return "N/A"
    valor = 0 if abs(valor) < 0.005 else valor
    return f"{valor:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")


def pdf_bytes(resultado, comparacao, indicador):
    stream = BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=landscape(A4), leftMargin=30, rightMargin=30,
                            topMargin=32, bottomMargin=48, title=resultado["nome"], author="Simulador Fundeb")
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="FundebTitle", fontName="FundebBold", fontSize=18, leading=23, spaceAfter=12, textColor=colors.HexColor("#204d46")))
    styles.add(ParagraphStyle(name="FundebBody", fontName="Fundeb", fontSize=9, leading=13, spaceAfter=7))
    styles.add(ParagraphStyle(name="FundebSmall", fontName="Fundeb", fontSize=7, leading=10))
    styles.add(ParagraphStyle(name="FundebRight", parent=styles["FundebSmall"], alignment=TA_RIGHT))
    p = lambda s, style="FundebBody": Paragraph(escape(str(s)), styles[style])
    story = [p("Simulador Fundeb | Relatório de cenário", "FundebTitle"), p(resultado["nome"]),
             p(f"Base: {resultado['base']['nome']} | Situação: {resultado['base']['status']}"),
             p(f"Exercício da base: {resultado['base'].get('ano_exercicio') or 'não identificado'} | Comparação: {comparacao}"),
             p(f"Cenário: {resultado['id']} | Calculado em: {resultado['criado_em']}"),
             p(f"Participação relativa: {resultado['recorte']['denominador']}."),
             p(resultado["aviso"]), p(resultado["nota_matriculas"], "FundebSmall")]
    entrada = resultado["entrada"].get("cenario", resultado["entrada"])
    receita = entrada["receita"]
    taxa = resultado.get("parametros_efetivos", {}).get("taxa_receita_percentual", Receita(**receita).taxa())
    rubricas = {"recursos_vaaf": "fundo antes da complementação VAAF", "recursos_vaat": "receita VAAT do snapshot",
                "outras_receitas_vaat": "outras receitas VAAT"}
    story.append(p(f"Receita nominal: {numero(taxa)}%{' ao ano' if resultado['tipo'] == 'trajetoria' else ''}. "
                   f"Rubricas: {', '.join(rubricas[r] for r in receita['rubricas'])}. "
                   f"Complementações: {receita['complementacoes']}. Fonte: {receita['fonte']}.", "FundebSmall"))
    if resultado["tipo"] == "cenario":
        params = resultado["parametros_efetivos"]["A_B"]
        story.append(p("Montantes A/B (R$): " + "; ".join(f"{m.upper()} {numero(params['complementacao_'+m])}" for m in ("vaaf", "vaat", "vaar")) +
                       f". NSE: {params['min_nse']} a {params['max_nse']}; fator fiscal: {params['min_nf']} a {params['max_nf']}.", "FundebSmall"))
        story.append(p(f"Ajustes: {len(resultado['ajustes'])} operações em {len({a['ibge'] for a in resultado['ajustes']})} redes. "
                       "A lista de alterações e os pesos efetivos constam do Excel e do resultado salvo.", "FundebSmall"))
    else:
        story.append(p("Hipótese da trajetória: " + resultado["entrada"]["hipotese"], "FundebSmall"))
    story.append(p("Identificação do conteúdo da base: " + resultado["base"]["fingerprint"], "FundebSmall"))
    story.append(p("Identificação do motor: " + resultado["motor_sha256"], "FundebSmall"))
    for aviso in resultado["avisos"]:
        story.append(p(aviso, "FundebSmall"))
    story.append(p("A participação usa recursos totais do Fundeb, independentemente do indicador escolhido. "
                   "N/A indica percentual ou valor por aluno sem denominador válido. "
                   "O Excel contém as quatro comparações e os dados nacionais completos dos cenários.", "FundebSmall"))
    conjuntos = [(None, resultado["comparativos"])] if resultado["tipo"] == "cenario" else [
        (s["ano"], s["comparativos"]) for s in resultado["series"]]
    for ano, comparativos in conjuntos:
        comp = comparativos[comparacao]
        story.extend([Spacer(1, 12), p((f"Ano hipotético {ano} | " if ano is not None else "") + METRICAS[indicador])])
        linhas = [[p(x, "FundebSmall") for x in ("UF", "Rede", "Referência", "Simulado", "Diferença", "Variação (%)", "Partic. (p.p.)")]]
        for row in comp["linhas"]:
            v = row["indicadores"][indicador]
            linhas.append([p(row["uf"], "FundebSmall"), p(row["nome"], "FundebSmall"),
                           *[p(numero(v[c]), "FundebRight") for c in ("base", "simulado", "delta", "percentual")],
                           p(numero(row["participacao"]["delta_pp"]), "FundebRight")])
        v = comp["totais"][indicador]
        linhas.append([p("", "FundebSmall"), p("Total do recorte", "FundebSmall"),
                       *[p(numero(v[c]), "FundebRight") for c in ("base", "simulado", "delta", "percentual")], p("", "FundebSmall")])
        tabela = LongTable(linhas, colWidths=[28, 183, 126, 126, 126, 90, 102], repeatRows=1)
        tabela.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e5eee9")),
                                   ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#f5f7f6")]),
                                   ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#e5eee9")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                                   ("TOPPADDING", (0, 0), (-1, -1), 3)]))
        story.append(tabela)

    def rodape(canvas, documento):
        canvas.setFont("Fundeb", 7)
        canvas.setFillColor(colors.HexColor("#53645e"))
        canvas.drawString(30, 25, "Cenário hipotético. Não constitui previsão de repasses. Base: " + resultado["base"]["id"])
        canvas.drawRightString(landscape(A4)[0] - 30, 25, f"Página {documento.page}")
    doc.build(story, onFirstPage=rodape, onLaterPages=rodape)
    return stream.getvalue()


def exportar(resultado, formato, comparacao="B-A", indicador="recursos_fundeb"):
    if comparacao not in {"B-A", "C-A", "D-A", "D-C"} or indicador not in METRICAS:
        raise ValueError("Comparação ou indicador desconhecido.")
    if formato == "csv":
        return csv_bytes(resultado, comparacao), "text/csv; charset=utf-8"
    if formato == "xlsx":
        return xlsx_bytes(resultado, comparacao), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if formato == "pdf":
        return pdf_bytes(resultado, comparacao, indicador), "application/pdf"
    raise ValueError("Formato deve ser pdf, xlsx ou csv.")
