"""
FND-08 — Exportações (CSV, XLSX, PDF) produzidas a partir do resultado já calculado.
Nenhuma exportação recalcula o cenário: todas leem o ResultadoCenario guardado.

CSV: UTF-8 com BOM, delimitador ';', separador decimal ','. As primeiras linhas,
iniciadas por '#', trazem os metadados (use comment='#' ao ler com pandas).
"""
from __future__ import annotations

import io
import json
import math

import pandas as pd

from app.services.cenarios import DESCRICAO_CENARIOS, ResultadoCenario
from app.services.comparacao import (
    DESCRICAO_RECORTES, EFEITOS, INDICADORES, efeitos_disponiveis, ibges_recorte,
    participacao_recorte, tabela_comparativa, totais_recorte,
)

CSV_SEPARADOR = ";"
CSV_DECIMAL = ","
_PREFIXOS_FORMULA = ("=", "+", "-", "@", "\t", "\r")


def neutralizar_formula(v):
    """Impede que textos sejam interpretados como fórmula por planilhas."""
    if isinstance(v, str) and v.startswith(_PREFIXOS_FORMULA):
        return "'" + v
    return v


def _texto_seguro(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    for col in df.columns:
        if df[col].dtype == object or pd.api.types.is_string_dtype(df[col]):
            df[col] = df[col].map(neutralizar_formula)
    return df


def nome_arquivo(res: ResultadoCenario, extensao: str) -> str:
    return f"cenario_{res.cenario_id[:8]}_{res.base['base_id']}.{extensao}"


def linhas_metadados(res: ResultadoCenario, recorte: str, n_linhas: int) -> list[tuple[str, str]]:
    b = res.base
    par = res.parametros_efetivos
    rec = res.receita
    linhas = [
        ("Cenário", res.cenario_id),
        ("Calculado em (UTC)", res.criado_em),
        ("Base", b["base_id"]),
        ("Exercício da base", str(b["ano_exercicio"]) if b["ano_exercicio"] else "não identificado"),
        ("Situação da base", f"{b['situacao']}{'' if b['homologada'] else ' (não homologada)'}"),
        ("Versão do motor", res.versao_motor),
        ("Recorte", f"{DESCRICAO_RECORTES[recorte]} — {n_linhas} redes"),
        ("Complementação VAAF (R$)", f"{par['complementacao_vaaf']:.2f}"),
        ("Complementação VAAT (R$)", f"{par['complementacao_vaat']:.2f}"),
        ("Complementação VAAR (R$)", f"{par['complementacao_vaar']:.2f}"),
        ("NSE (mín–máx)", f"{par['min_nse']}–{par['max_nse']}"),
        ("NF (mín–máx)", f"{par['min_nf']}–{par['max_nf']}"),
        ("Pesos", "alterados pelo usuário" if par["pesos_vaaf"] or par["pesos_vaat"] else "da base"),
        ("Receita", "constante" if rec["tipo"] == "constante" else
         f"crescimento de {rec['taxa'] * 100:.4g}% em {', '.join(rec['rubricas'])}; "
         f"complementações {rec['complementacoes'].replace('_', ' ')}"
         + (f"; fonte: {rec['fonte_taxa']}" if rec.get("fonte_taxa") else "")),
        ("Cenários", "; ".join(f"{k}: {v}" for k, v in DESCRICAO_CENARIOS.items() if k in res.cenarios)),
        ("Ajustes", str(res.resumo_ajustes["quantidade"])),
        ("Validação interna", "; ".join(f"{k}: {'ok' if v['valido'] else 'com erros'}" for k, v in res.validacoes.items())),
        ("Hashes da base", json.dumps(b["hashes"], ensure_ascii=False)),
        ("Aviso", res.aviso),
    ]
    return linhas


def _comparativo_exportavel(res: ResultadoCenario, recorte: str, selecionadas) -> pd.DataFrame:
    tab = tabela_comparativa(res, recorte, selecionadas)
    tab.insert(0, "cenario_id", res.cenario_id)
    tab.insert(1, "base_id", res.base["base_id"])
    return tab


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------

def exportar_csv(res: ResultadoCenario, recorte: str = "estaduais_df", selecionadas=None) -> bytes:
    tab = _comparativo_exportavel(res, recorte, selecionadas)
    buf = io.StringIO()
    for chave, valor in linhas_metadados(res, recorte, len(tab)):
        buf.write(f"# {chave}: {valor}\n".replace("\r", " "))
    _texto_seguro(tab).to_csv(buf, sep=CSV_SEPARADOR, decimal=CSV_DECIMAL, index=False, lineterminator="\n")
    return buf.getvalue().encode("utf-8-sig")


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------

def _resumo_df(res: ResultadoCenario, ids: list[int]) -> pd.DataFrame:
    tot = totais_recorte(res, ids)
    linhas = []
    for col, rotulo, unidade, agreg in INDICADORES:
        linha = {"Indicador": rotulo, "Unidade": unidade,
                 "Agregação": "soma" if agreg == "soma" else "razão de somas"}
        for c in res.cenarios:
            linha[f"Cenário {c}"] = tot["por_cenario"][c][col]
        for ef, info in tot["efeitos"].items():
            linha[f"{ef.replace('_', ' − ')} (dif.)"] = info["indicadores"][col]["diferenca"]
            linha[f"{ef.replace('_', ' − ')} (%)"] = info["indicadores"][col]["variacao_pct"]
        linhas.append(linha)
    part = participacao_recorte(res, ids)
    linha = {"Indicador": "Participação do recorte nos recursos nacionais do Fundeb", "Unidade": "%",
             "Agregação": part["denominador"]}
    for c in res.cenarios:
        linha[f"Cenário {c}"] = part["participacao_pct"][c]
    for ef, v in part["variacao_pp"].items():
        linha[f"{ef.replace('_', ' − ')} (dif.)"] = v
    linhas.append(linha)
    return pd.DataFrame(linhas)


def _dados_completos(res: ResultadoCenario) -> pd.DataFrame:
    partes = []
    tipos = res.entes.set_index("ibge")["tipo_rede"]
    for c in res.cenarios:
        df = res.resultados[c].copy()
        df.insert(0, "cenario", c)
        df.insert(4, "tipo_rede", df["ibge"].map(tipos))
        partes.append(df)
    return pd.concat(partes, ignore_index=True)


def exportar_xlsx(res: ResultadoCenario, recorte: str = "estaduais_df", selecionadas=None) -> bytes:
    from openpyxl.utils import get_column_letter

    ids = ibges_recorte(res, recorte, selecionadas)
    comparativo = _comparativo_exportavel(res, recorte, selecionadas)
    folhas = {
        "Resumo": _resumo_df(res, ids),
        "Comparativo": comparativo,
        "Dados completos": _dados_completos(res),
        "Ajustes": res.ajustes if len(res.ajustes) else pd.DataFrame({"mensagem": ["Cenário sem ajustes de matrículas"]}),
        "Metadados": pd.DataFrame(linhas_metadados(res, recorte, len(comparativo)), columns=["Campo", "Valor"]),
    }
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for nome, df in folhas.items():
            df.to_excel(writer, sheet_name=nome, index=False)
            ws = writer.sheets[nome]
            ws.auto_filter.ref = ws.dimensions
            ws.freeze_panes = "A2"
            for cells in ws.iter_rows(min_row=2):
                for cell in cells:
                    # Textos são gravados como texto, mesmo que comecem com '='.
                    if isinstance(cell.value, str) and cell.data_type == "f":
                        cell.data_type = "s"
                    if isinstance(cell.value, float) and not math.isfinite(cell.value):
                        cell.value = None
            for i, col in enumerate(df.columns, start=1):
                largura = min(max(len(str(col)), 10) + 2, 60)
                ws.column_dimensions[get_column_letter(i)].width = largura
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _fmt_num(v, casas=0):
    if v is None or (isinstance(v, float) and not math.isfinite(v)):
        return "n/a"
    s = f"{v:,.{casas}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def exportar_pdf(res: ResultadoCenario, recorte: str = "estaduais_df", selecionadas=None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ids = ibges_recorte(res, recorte, selecionadas)
    tab = tabela_comparativa(res, recorte, selecionadas)
    estilos = getSampleStyleSheet()
    pequeno = estilos["BodyText"].clone("pequeno", fontSize=8, leading=10)
    cabecalho = estilos["BodyText"].clone("cabecalho", fontName="Helvetica-Bold", fontSize=7, leading=8)

    def cab_(textos):
        return [Paragraph(t, cabecalho) for t in textos]
    buf = io.BytesIO()

    def rodape(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.drawString(12 * mm, 8 * mm, res.aviso)
        canvas.drawRightString(doc.pagesize[0] - 12 * mm, 8 * mm,
                               f"Cenário {res.cenario_id[:8]} · base {res.base['base_id']} · página {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=16 * mm,
                            title=f"Cenário {res.cenario_id[:8]} — Simulador Fundeb")
    story = [Paragraph("Simulador Fundeb — comparação de cenários", estilos["Title"])]
    if not res.base["homologada"]:
        story.append(Paragraph("<b>Base não homologada.</b> Resultados não podem ser apresentados como "
                               "referentes a um exercício oficial.", pequeno))
    meta = [[Paragraph(f"<b>{k}</b>", pequeno), Paragraph(str(v), pequeno)]
            for k, v in linhas_metadados(res, recorte, len(tab)) if k != "Hashes da base"]
    t = Table(meta, colWidths=[55 * mm, 210 * mm])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.lightgrey)]))
    story += [Spacer(1, 4 * mm), t, Spacer(1, 6 * mm)]

    # Resumo do recorte
    tot = totais_recorte(res, ids)
    cab = ["Indicador", "Unidade"] + [f"Cenário {c}" for c in res.cenarios] + \
          [f"{EFEITOS[e][0]} − {EFEITOS[e][1]}" for e in efeitos_disponiveis(res)]
    linhas = [cab_(cab)]
    for col, rotulo, unidade, agreg in INDICADORES:
        casas = 2 if agreg == "razao" else 0
        linha = [rotulo, unidade] + [_fmt_num(tot["por_cenario"][c][col], casas) for c in res.cenarios]
        linha += [_fmt_num(tot["efeitos"][e]["indicadores"][col]["diferenca"], casas) for e in efeitos_disponiveis(res)]
        linhas.append(linha)
    part = participacao_recorte(res, ids)
    linhas.append(["Participação nos recursos nacionais", "%"]
                  + [_fmt_num(part["participacao_pct"][c], 4) for c in res.cenarios]
                  + [_fmt_num(part["variacao_pp"][e], 4) + " p.p." for e in efeitos_disponiveis(res)])
    story.append(Paragraph(f"Resumo — {DESCRICAO_RECORTES[recorte]} ({len(ids)} redes). "
                           "Valores por aluno nos totais: razão de somas.", estilos["Heading3"]))
    t = Table(linhas, repeatRows=1)
    t.setStyle(_estilo_tabela(colors))
    story += [t, Spacer(1, 6 * mm)]

    # Tabela comparativa paginada (A × B)
    story.append(Paragraph("Comparativo por rede — referência (A) × simulado (B)", estilos["Heading3"]))
    cab = ["UF", "Rede", "Matrículas A", "Matrículas B", "Recursos Fundeb A (R$)", "Recursos Fundeb B (R$)",
           "Variação (R$)", "Variação (%)", "Compl. VAAF B−A (R$)", "Compl. VAAT B−A (R$)",
           "VAAF/aluno A→B", "VAAT/aluno A→B"]
    linhas = [cab_(cab)]
    for r in tab.itertuples():
        linhas.append([
            r.uf, Paragraph(str(r.nome), pequeno),
            _fmt_num(r.matriculas_total_A), _fmt_num(r.matriculas_total_B),
            _fmt_num(r.recursos_fundeb_A), _fmt_num(r.recursos_fundeb_B),
            _fmt_num(r.recursos_fundeb_dif_B_A), _fmt_num(r.recursos_fundeb_pct_B_A, 2),
            _fmt_num(r.complemento_vaaf_dif_B_A), _fmt_num(r.complemento_vaat_dif_B_A),
            f"{_fmt_num(r.vaaf_final_A, 2)} → {_fmt_num(r.vaaf_final_B, 2)}",
            f"{_fmt_num(r.vaat_final_A, 2)} → {_fmt_num(r.vaat_final_B, 2)}",
        ])
    larguras = [10, 38, 20, 20, 28, 28, 24, 16, 24, 24, 28, 28]
    t = Table(linhas, repeatRows=1, colWidths=[w * mm for w in larguras])
    t.setStyle(_estilo_tabela(colors))
    story.append(t)
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph("n/a: não aplicável (denominador zero ou valor ausente). "
                           "Valores nominais, em reais; VAAF/VAAT por aluno em R$/aluno/ano.", pequeno))
    doc.build(story, onFirstPage=rodape, onLaterPages=rodape)
    return buf.getvalue()


def _estilo_tabela(colors):
    from reportlab.platypus import TableStyle
    return TableStyle([
        ("FONT", (0, 0), (-1, -1), "Helvetica", 7),
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 7),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8E2D6")),
        ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F5F0")]),
    ])


EXPORTADORES = {
    "csv": (exportar_csv, "text/csv; charset=utf-8"),
    "xlsx": (exportar_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "pdf": (exportar_pdf, "application/pdf"),
}
