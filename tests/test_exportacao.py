"""FND-08 — exportações do mesmo cenário calculado."""
import io

import pandas as pd
import pytest

from schemas.cenarios import AVISO_METODOLOGICO, CenarioRequest
from services.cenarios import executar_cenario
from services.exportacao import exportar_csv, exportar_pdf, exportar_xlsx, neutralizar_formula
from services.comparacao import tabela_comparativa


@pytest.fixture
def resultado(base_sintetica):
    base_sintetica.complementar.loc[base_sintetica.complementar.ibge == 23, "nome"] = "=HYPERLINK(\"x\")"
    base_sintetica.entes.loc[base_sintetica.entes.ibge == 23, "nome"] = "=HYPERLINK(\"x\")"
    return executar_cenario(CenarioRequest(
        ajustes=[{"ibge": 22, "categoria": "fundamental", "operacao": "acrescentar", "valor": 10}],
        receita={"tipo": "crescimento", "taxa": 0.03},
    ), base_sintetica)


def test_csv_tem_todas_as_linhas_metadados_e_valores(resultado):
    conteudo = exportar_csv(resultado)
    assert conteudo.startswith("﻿".encode("utf-8"))
    texto = conteudo.decode("utf-8-sig")
    assert AVISO_METODOLOGICO in texto
    assert resultado.cenario_id in texto
    df = pd.read_csv(io.StringIO(texto), sep=";", decimal=",", comment="#")
    esperado = tabela_comparativa(resultado)
    assert len(df) == len(esperado)
    assert df["recursos_fundeb_B"].sum() == pytest.approx(esperado["recursos_fundeb_B"].sum(), abs=0.01)
    assert (df["base_id"] == "sintetica").all()
    assert "'=HYPERLINK" in texto


def test_xlsx_abas_e_valores(resultado):
    from openpyxl import load_workbook
    conteudo = exportar_xlsx(resultado)
    wb = load_workbook(io.BytesIO(conteudo))
    assert wb.sheetnames == ["Resumo", "Comparativo", "Dados completos", "Ajustes", "Metadados"]
    comp = pd.read_excel(io.BytesIO(conteudo), sheet_name="Comparativo")
    esperado = tabela_comparativa(resultado)
    assert len(comp) == len(esperado)
    assert comp["recursos_fundeb_D"].sum() == pytest.approx(esperado["recursos_fundeb_D"].sum())
    completos = pd.read_excel(io.BytesIO(conteudo), sheet_name="Dados completos")
    assert len(completos) == 4 * len(resultado.resultados["A"])  # cenários A–D, sem limite de linhas
    # Texto iniciado por '=' preservado como texto, não como fórmula
    ws = wb["Comparativo"]
    nomes = [c.value for c in ws["F"]] + [c.value for c in ws["E"]]
    celulas = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and "HYPERLINK" in c.value]
    assert celulas and all(c.data_type == "s" for c in celulas), nomes
    meta = pd.read_excel(io.BytesIO(conteudo), sheet_name="Metadados")
    assert AVISO_METODOLOGICO in meta["Valor"].tolist()


def test_pdf_gerado(resultado):
    conteudo = exportar_pdf(resultado)
    assert conteudo.startswith(b"%PDF")
    assert len(conteudo) > 2000


def test_exportacao_usa_resultado_guardado(resultado, base_sintetica):
    """Alterar a base depois do cálculo não altera a exportação do cenário já calculado."""
    antes = exportar_csv(resultado)
    base_sintetica.matriculas["fundamental"] = 0.0
    base_sintetica.complementar["recursos_vaaf"] = 1.0
    assert exportar_csv(resultado) == antes


@pytest.mark.parametrize("valor, esperado", [
    ("=1+1", "'=1+1"), ("+SOMA(A1)", "'+SOMA(A1)"), ("@x", "'@x"), ("Ceará", "Ceará"), (10.5, 10.5),
])
def test_neutralizar_formula(valor, esperado):
    assert neutralizar_formula(valor) == esperado
