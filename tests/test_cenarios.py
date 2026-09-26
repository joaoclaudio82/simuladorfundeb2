from copy import deepcopy
from io import BytesIO
import csv
import json

import numpy as np
import pandas as pd
import pytest
from pydantic import ValidationError
from openpyxl import load_workbook
from pypdf import PdfReader

from schemas.cenarios import CenarioRequest, Receita, TrajetoriaRequest
from services.bases import BaseRepository
from services.cenarios import aplicar_ajustes, fatores_amazonicos
from services.exportacao import exportar
from services.storage import ScenarioStore
from simulador import simula_fundeb, reescala_vetor


def request(**kwargs):
    return CenarioRequest(base_id="teste", **kwargs)


def test_zero_ajustes_e_taxa_zero_reproduzem_mesmo_cenario(servico):
    r = servico.executar(request())
    assert r["dados"]["A"] == r["dados"]["B"] == r["dados"]["C"] == r["dados"]["D"]
    for comp in r["comparativos"].values():
        assert comp["totais"]["recursos_fundeb"]["delta"] == 0
    assert {x["uf"] for x in r["comparativos"]["B-A"]["linhas"]} == {"CE", "PI", "AP", "DF"}


def test_conjunto_equivale_a_execucao_nacional_e_nao_muta_base(servico, base_teste):
    original = base_teste.tabelas["matriculas"].copy(deep=True)
    ajustes = [{"ibge": 23, "etapa": "tecnico", "operacao": "adicionar", "valor": 30},
               {"ibge": 22, "etapa": "regular", "operacao": "adicionar", "valor": 40}]
    r = servico.executar(request(ajustes=ajustes))
    direto = original.copy()
    direto.loc[direto.ibge == 23, "tecnico"] += 30
    direto.loc[direto.ibge == 22, "regular"] += 40
    esperado = simula_fundeb(direto, base_teste.tabelas["complementar"], base_teste.tabelas["pesos"],
                             **base_teste.meta["parametros"], arredondar=False).set_index("ibge")
    recebido = pd.DataFrame(r["dados"]["B"]).set_index("ibge")
    np.testing.assert_allclose(recebido.loc[esperado.index, "recursos_fundeb"], esperado.recursos_fundeb)
    pd.testing.assert_frame_equal(original, base_teste.tabelas["matriculas"])
    invertido = servico.executar(request(ajustes=list(reversed(ajustes))))
    assert r["dados"] == invertido["dados"]
    # Uma rede sem ajuste também participa da redistribuição.
    a = pd.DataFrame(r["dados"]["A"]).set_index("ibge")
    assert recebido.loc[2300011, "recursos_fundeb"] != a.loc[2300011, "recursos_fundeb"]


def test_conversao_conserva_total_e_percentual_arredonda_incremento(base_teste):
    req = request(ajustes=[{"ibge": 23, "etapa": "regular", "operacao": "converter", "destino": "tecnico", "valor": 50},
                           {"ibge": 16, "etapa": "tecnico", "operacao": "percentual", "valor": 10}])
    m, log = aplicar_ajustes(base_teste, req.ajustes)
    assert m.loc[m.ibge == 23, ["regular", "tecnico"]].sum(axis=1).item() == 120
    assert m.loc[m.ibge == 16, "tecnico"].item() == 6  # 0,5 unidade -> 1
    assert log[1]["delta_total"] == 0  # log ordenado pelo IBGE: AP antes de CE


@pytest.mark.parametrize("ajustes", [
    [{"ibge": 999, "etapa": "regular", "valor": 1}],
    [{"ibge": 23, "etapa": "ausente", "valor": 1}],
    [{"ibge": 23, "etapa": "regular", "valor": 500, "operacao": "converter", "destino": "tecnico"}],
    [{"ibge": 23, "etapa": "regular", "valor": 5}, {"ibge": 23, "etapa": "regular", "valor": 5}],
    [{"ibge": 23, "etapa": "regular", "valor": 5, "operacao": "converter", "destino": "tecnico"}, {"ibge": 23, "etapa": "tecnico", "valor": 5}],
])
def test_ajustes_invalidos_nao_sao_ignorados(base_teste, ajustes):
    with pytest.raises(ValueError):
        aplicar_ajustes(base_teste, request(ajustes=ajustes).ajustes)


@pytest.mark.parametrize("valor", [-1, 2.5, float("nan"), float("inf")])
def test_rejeita_contagens_invalidas(valor):
    with pytest.raises(ValidationError):
        request(ajustes=[{"ibge": 23, "etapa": "regular", "valor": valor}])


def test_receita_quatro_cenarios_montantes_e_rubricas(servico):
    r = servico.executar(request(receita={"taxa_percentual": 10, "rubricas": ["recursos_vaaf"], "complementacoes": "proporcionais"},
                                  ajustes=[{"ibge": 23, "etapa": "tecnico", "valor": 30, "operacao": "adicionar"}]))
    a, c = [pd.DataFrame(r["dados"][k]).set_index("ibge") for k in ("A", "C")]
    assert c.recursos_vaaf.sum() == pytest.approx(a.recursos_vaaf.sum() * 1.1)
    np.testing.assert_allclose(a.recursos_vaat, c.recursos_vaat)
    assert c.complemento_vaaf.sum() == pytest.approx(11000)
    assert c.complemento_vaat.sum() == pytest.approx(8800)
    totals = r["comparativos"]
    assert totals["D-A"]["totais"]["recursos_fundeb"]["delta"] == pytest.approx(
        totals["C-A"]["totais"]["recursos_fundeb"]["delta"] + totals["D-C"]["totais"]["recursos_fundeb"]["delta"])


def test_pib_calcula_cagr_e_exige_fonte():
    r = Receita(metodo="pib_cagr", pib=[{"ano": 2020, "valor": 100}, {"ano": 2022, "valor": 121}], fonte="Série sintética de teste")
    assert r.taxa() == pytest.approx(10)
    with pytest.raises(ValidationError):
        Receita(metodo="pib_cagr", pib=r.pib)
    with pytest.raises(ValidationError):
        Receita(metodo="pib_cagr", pib=[r.pib[0], r.pib[0]], fonte="Série sintética")


def test_recortes_inabilitados_e_denominador_nacional(servico):
    todos = servico.executar(request(recorte={"tipo": "todas"}))
    assert len(todos["comparativos"]["B-A"]["linhas"]) == 6  # inclui o inabilitado VAAT
    filtrado = servico.executar(request(recorte={"tipo": "estaduais", "ufs": ["CE"]}))
    row = filtrado["comparativos"]["B-A"]["linhas"][0]
    assert 0 < row["participacao"]["base_pct"] < 100
    assert row["indicadores"]["complemento_vaar"]["percentual"] is None
    with pytest.raises(ValueError, match="Propag"):
        servico.executar(request(recorte={"tipo": "propag"}))
    with pytest.raises(ValueError, match="UF"):
        servico.executar(request(recorte={"ufs": ["XX"]}))


def test_modelo_por_componentes_recompoe_vaat(base_teste):
    b = deepcopy(base_teste)
    comp = b.tabelas["complementar"]
    comp["outras_receitas_vaat"] = 1000.
    p = b.meta["parametros"]
    r = simula_fundeb(b.tabelas["matriculas"], comp, b.tabelas["pesos"], **p, modo_vaat="componentes", arredondar=False)
    np.testing.assert_allclose(r.recursos_vaat, r.recursos_vaaf_final + 1000)
    with pytest.raises(ValueError, match="explícitas"):
        simula_fundeb(b.tabelas["matriculas"], comp.drop(columns="outras_receitas_vaat"), b.tabelas["pesos"], **p, modo_vaat="componentes")


def test_fator_amazonico_exige_regra_e_vigencia(base_teste):
    with pytest.raises(ValueError, match="norma"):
        fatores_amazonicos(base_teste, True)
    base_teste.meta["amazonico"] = {"status": "validado", "norma": "Regra SINTÉTICA para teste", "responsavel_validacao": "teste",
        "ano_inicio": 2026, "ano_fim": 2026, "incidencia": "matriculas_ponderadas_por_ente"}
    f = base_teste.tabelas["entes"][["ibge"]].copy()
    f["fator_vaaf"] = 1.
    f["fator_vaat"] = 1.
    f.loc[f.ibge == 16, ["fator_vaaf", "fator_vaat"]] = 1.2
    base_teste.tabelas["fatores_amazonicos"] = f
    fatores = fatores_amazonicos(base_teste, True)
    args = [base_teste.tabelas[k] for k in ("matriculas", "complementar", "pesos")]
    a = simula_fundeb(*args, **base_teste.meta["parametros"], arredondar=False)
    b = simula_fundeb(*args, **base_teste.meta["parametros"], fatores_matriculas=fatores, arredondar=False)
    assert b.set_index('ibge').loc[16, 'matriculas_vaaf'] == pytest.approx(a.set_index('ibge').loc[16, 'matriculas_vaaf'] * 1.2)
    base_teste.meta["ano_exercicio"] = 2025
    with pytest.raises(ValueError, match="vigência"):
        fatores_amazonicos(base_teste, True)


def test_trajetoria_linear_base_fixa_receita_composta(servico):
    r = servico.trajetoria(TrajetoriaRequest(cenario=request(receita={"taxa_percentual": 10}), ano_inicial=2026, ano_final=2028,
                entes=[23], etapa="tecnico", aumento_total_percentual=50, hipotese="Hipótese sintética de teste"))
    assert [x["incremento_ept"] for x in r["series"]] == [0, 5, 10]
    assert [x["taxa_receita_acumulada"] for x in r["series"]] == pytest.approx([0, 10, 21])
    assert r["series"][0]["comparativos"]["D-A"]["totais"]["recursos_fundeb"]["delta"] == 0
    assert servico.store.obter(r["id"])["tipo"] == "trajetoria"


def test_snapshot_persiste_e_nao_recalcula_com_base_alterada(servico, tmp_path):
    r = servico.executar(request())
    reaberto = ScenarioStore(tmp_path / "cenarios.sqlite3").obter(r["id"])
    assert reaberto == r
    servico.bases.obter = lambda _: (_ for _ in ()).throw(AssertionError("Não deveria consultar a base"))
    for formato in ("csv", "xlsx", "pdf"):
        data, _ = exportar(reaberto, formato)
        assert len(data) > 100
    with pytest.raises(KeyError):
        servico.store.obter("../cenarios.sqlite3")


def test_exportacoes_valores_completos_texto_seguro_e_metadados(servico):
    r = servico.executar(request(ajustes=[{"ibge": 23, "etapa": "tecnico", "operacao": "adicionar", "valor": 10}]))
    r["comparativos"]["B-A"]["linhas"][0]["nome"] = '=HYPERLINK("https://example.invalid")'
    csv_data, _ = exportar(r, "csv")
    rows = list(csv.DictReader(csv_data.decode("utf-8-sig").splitlines(), delimiter=";"))
    assert len(rows) == 4
    assert rows[0]["nome"].startswith("'=")
    assert rows[0]["base_sha256"] == r["base"]["fingerprint"]
    assert float(rows[0]["recursos_fundeb_simulado"]) == r["comparativos"]["B-A"]["linhas"][0]["indicadores"]["recursos_fundeb"]["simulado"]
    xlsx, _ = exportar(r, "xlsx")
    wb = load_workbook(BytesIO(xlsx), read_only=True, data_only=False)
    assert wb["Dados A"].max_row == 7
    headers = [c.value for c in wb["Comparativo"][1]]
    assert wb["Comparativo"].cell(2, headers.index("nome") + 1).data_type == "s"
    assert "Metadados" in wb.sheetnames and "Ajustes" in wb.sheetnames
    pdf, _ = exportar(r, "pdf")
    text = "".join(p.extract_text() for p in PdfReader(BytesIO(pdf)).pages)
    assert "Ceará" in text and "Distrito Federal" in text and r["id"] in text


def test_base_real_cobertura_fracoes_e_exercicio_desconhecido():
    b = BaseRepository().obter("legado")
    assert b.resumo()["redes"] == 5595
    assert b.resumo()["tipos"] == {"municipal": 5568, "estadual": 26, "distrital": 1}
    assert b.meta["ano_exercicio"] is None
    assert ((b.tabelas["matriculas"][b.etapas] % 1) != 0).sum().sum() == 804


def test_vetor_constante_respeita_intervalo_configurado():
    np.testing.assert_allclose(reescala_vetor(np.ones(3), maximo=2, minimo=2), [2, 2, 2])


def test_zero_matriculas_e_pesos_vaar_invalidos(base_teste):
    b = base_teste
    mat = b.tabelas["matriculas"].copy()
    mat.loc[mat.ibge == 2300011, ["regular", "tecnico"]] = 0
    r = simula_fundeb(mat, b.tabelas["complementar"], b.tabelas["pesos"], **b.meta["parametros"], arredondar=False)
    assert np.isnan(r.set_index("ibge").loc[2300011, "vaaf_final"])
    mat.loc[mat.ibge == 23, ["regular", "tecnico"]] = 0
    with pytest.raises(ValueError, match="UF"):
        simula_fundeb(mat, b.tabelas["complementar"], b.tabelas["pesos"], **b.meta["parametros"])
    comp = b.tabelas["complementar"].copy()
    comp["peso_vaar"] = 0
    with pytest.raises(ValueError, match="VAAR"):
        simula_fundeb(b.tabelas["matriculas"], comp, b.tabelas["pesos"],
                       **{**b.meta["parametros"], "complementacao_vaar": 10})
